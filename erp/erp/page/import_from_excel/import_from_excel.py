import frappe
from frappe import _


@frappe.whitelist()
def get_sheet_names(file_url: str) -> list:
	"""Return all sheet names from an Excel file."""
	import openpyxl

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	wb = openpyxl.load_workbook(file_doc.get_full_path(), read_only=True)
	names = wb.sheetnames
	wb.close()
	return names


@frappe.whitelist()
def parse_excel(file_url: str, sheet_name: str = None) -> dict:
	"""Parse uploaded Excel file. Returns column metadata and sample values."""
	import openpyxl

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	wb = openpyxl.load_workbook(file_doc.get_full_path(), read_only=True, data_only=True)

	if sheet_name and sheet_name in wb.sheetnames:
		ws = wb[sheet_name]
	else:
		ws = wb.active

	all_rows = []
	for row in ws.iter_rows(values_only=True):
		all_rows.append(row)
		if len(all_rows) > 101:
			break

	wb.close()

	if not all_rows:
		return {"columns": [], "total_rows": 0}

	headers = [
		str(h).strip() if h is not None else f"Column {i + 1}"
		for i, h in enumerate(all_rows[0])
	]
	data_rows = all_rows[1:]

	columns = []
	for i, name in enumerate(headers):
		samples = []
		for row in data_rows[:5]:
			val = row[i] if i < len(row) else None
			samples.append(str(val) if val is not None else "")

		col_type = _infer_col_type(data_rows, i)
		columns.append(
			{
				"id": f"x{i + 1}",
				"name": name,
				"sample": samples[0] if samples else "",
				"type": col_type,
			}
		)

	return {"columns": columns, "total_rows": len(data_rows)}


def _infer_col_type(rows: list, col_idx: int) -> str:
	import datetime

	for row in rows[:20]:
		if col_idx < len(row) and row[col_idx] is not None:
			v = row[col_idx]
			if isinstance(v, (int, float)):
				return "number"
			if isinstance(v, (datetime.date, datetime.datetime)):
				return "date"
	return "text"


@frappe.whitelist()
def get_doctype_fields(doctype: str) -> dict:
	"""Return importable fields of a DocType with metadata."""
	SKIP_TYPES = {
		"Tab Break",
		"Section Break",
		"Column Break",
		"HTML",
		"Button",
		"Fold",
		"Heading",
		"HTML Editor",
		"Markdown Editor",
	}

	meta = frappe.get_meta(doctype)
	fields = []
	for f in meta.fields:
		if f.fieldtype in SKIP_TYPES:
			continue
		fields.append(
			{
				"id": f"f_{f.fieldname}",
				"fieldname": f.fieldname,
				"label": f.label or f.fieldname,
				"fieldtype": f.fieldtype,
				"reqd": bool(f.reqd),
				"options": f.options or "",
			}
		)

	return {"fields": fields, "doctype": doctype}


@frappe.whitelist()
def get_all_doctypes() -> list:
	"""Return all non-single, non-child DocTypes in the erp module."""
	return frappe.get_all(
		"DocType",
		filters={"module": "erp", "issingle": 0, "istable": 0},
		fields=["name"],
		order_by="name asc",
		pluck="name",
	)


@frappe.whitelist()
def get_preview_data(doctype: str, mappings: str, file_url: str, sheet_name: str = None) -> dict:
	"""Return first 10 mapped rows for the validate step."""
	import json
	import openpyxl

	mapping_list = json.loads(mappings) if isinstance(mappings, str) else mappings

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	wb = openpyxl.load_workbook(file_doc.get_full_path(), read_only=True, data_only=True)

	ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
	all_rows = list(ws.iter_rows(values_only=True))
	wb.close()

	if not all_rows:
		return {"headers": [], "rows": [], "total": 0}

	raw_headers = [
		str(h).strip() if h is not None else f"Column {i + 1}"
		for i, h in enumerate(all_rows[0])
	]

	col_map: dict[int, str] = {}
	for m in mapping_list:
		ec = m.get("excel_col")
		fn = m.get("fieldname")
		if ec in raw_headers:
			col_map[raw_headers.index(ec)] = fn

	out_headers = [m["fieldname"] for m in mapping_list if m.get("excel_col") in raw_headers]
	out_rows = []

	for row in all_rows[1:11]:
		out_row = {}
		for col_idx, fieldname in col_map.items():
			val = row[col_idx] if col_idx < len(row) else None
			out_row[fieldname] = str(val) if val is not None else ""
		out_rows.append(out_row)

	return {"headers": out_headers, "rows": out_rows, "total": len(all_rows) - 1}


@frappe.whitelist(methods=["POST"])
def import_records(
	doctype: str, mappings: str, file_url: str, sheet_name: str = None
) -> dict:
	"""Import all rows from Excel using the provided column → field mapping."""
	import json
	import openpyxl

	frappe.has_permission(doctype, ptype="create", throw=True)

	mapping_list = json.loads(mappings) if isinstance(mappings, str) else mappings

	file_doc = frappe.get_doc("File", {"file_url": file_url})
	wb = openpyxl.load_workbook(file_doc.get_full_path(), read_only=True, data_only=True)

	ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
	all_rows = list(ws.iter_rows(values_only=True))
	wb.close()

	if not all_rows:
		return {"imported": 0, "skipped": 0, "errors": []}

	raw_headers = [
		str(h).strip() if h is not None else f"Column {i + 1}"
		for i, h in enumerate(all_rows[0])
	]

	col_map: dict[int, str] = {}
	for m in mapping_list:
		ec = m.get("excel_col")
		fn = m.get("fieldname")
		if ec in raw_headers:
			col_map[raw_headers.index(ec)] = fn

	meta = frappe.get_meta(doctype)
	NUMERIC_TYPES = {"Int", "Float", "Currency", "Percent"}

	imported = 0
	skipped = 0
	errors: list[dict] = []

	for row_num, row in enumerate(all_rows[1:], 2):
		doc_data: dict = {"doctype": doctype}
		for col_idx, fieldname in col_map.items():
			val = row[col_idx] if col_idx < len(row) else None
			if val is None:
				continue
			field_meta = meta.get_field(fieldname)
			if field_meta and field_meta.fieldtype in NUMERIC_TYPES:
				doc_data[fieldname] = val
			else:
				str_val = str(val).strip()
				if str_val:
					doc_data[fieldname] = str_val

		try:
			doc = frappe.get_doc(doc_data)
			doc.insert(ignore_permissions=True)
			imported += 1
		except Exception as e:
			skipped += 1
			if len(errors) < 100:
				errors.append({"row": row_num, "error": str(e)[:400]})

	frappe.db.commit()
	return {"imported": imported, "skipped": skipped, "errors": errors}


@frappe.whitelist()
def check_can_auto_create(doctype: str) -> dict:
	"""Check whether missing linked records can be auto-created from just their name."""
	SKIP_TYPES = {
		"Tab Break", "Section Break", "Column Break", "HTML",
		"Button", "Fold", "Heading", "HTML Editor", "Markdown Editor",
	}

	meta = frappe.get_meta(doctype)

	required_fields = [
		f.fieldname
		for f in meta.fields
		if f.reqd and f.fieldtype not in SKIP_TYPES
	]

	autoname_field = None
	if meta.autoname and meta.autoname.startswith("field:"):
		autoname_field = meta.autoname.replace("field:", "")

	# Auto-create is safe only when the single required field IS the autoname field
	can_auto = (
		autoname_field is not None
		and len(required_fields) == 1
		and required_fields[0] == autoname_field
	)

	return {
		"can_auto": can_auto,
		"autoname_field": autoname_field,
		"required_fields": required_fields,
	}


@frappe.whitelist(methods=["POST"])
def auto_create_linked(doctype: str, values: str) -> dict:
	"""Create minimal records for missing linked values (name-only doctypes)."""
	import json

	frappe.has_permission(doctype, ptype="create", throw=True)

	values_list = json.loads(values) if isinstance(values, str) else values

	meta = frappe.get_meta(doctype)
	if not meta.autoname or not meta.autoname.startswith("field:"):
		frappe.throw(_("Cannot auto-create records for this DocType"))

	name_field = meta.autoname.replace("field:", "")

	created: list[str] = []
	already_existed: list[str] = []
	errors: list[dict] = []

	for val in values_list:
		if frappe.db.exists(doctype, val):
			already_existed.append(val)
			continue
		try:
			doc = frappe.get_doc({"doctype": doctype, name_field: val})
			doc.insert(ignore_permissions=True)
			created.append(val)
		except Exception as e:
			errors.append({"value": val, "error": str(e)[:200]})

	frappe.db.commit()
	return {"created": created, "already_existed": already_existed, "errors": errors}
